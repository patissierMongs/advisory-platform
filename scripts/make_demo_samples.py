#!/usr/bin/env python3
"""데모 샘플 생성기 — samples/demo_advisories/ 를 통째로 재생성한다.

폴더째 업로드·출처 자동 탐지·CVE 게이트 데모를 폐쇄망에서 바로 돌릴 수 있게,
권고문 PDF(폴더 구조 포함)와 그 안의 CVE 코드를 전부 커버하는 내부 형식
CVE 피드(cve_feed_demo.json)를 한 묶음으로 만든다.

  python scripts/make_demo_samples.py            # samples/demo_advisories/ 생성
  python scripts/make_demo_samples.py --zip PATH # + 전체를 zip 으로 묶기

데모 시나리오(출처 탐지 우선순위 — 상위 폴더명 → 폴더명 → 파일명 → 본문):
  국가정보원/2026-07/…          → 상위 폴더명에서 '국가정보원' 탐지
  KISA/…                        → 폴더명에서 'KISA' 탐지
  NCSC/웹브라우저/…             → 상위 폴더명에서 'NCSC' 탐지
  FSEC-2026-0021 ….pdf          → 파일명 별칭(FSEC)에서 '금융보안원' 탐지
  기타접수/advisory-unknown-001 → 본문 'KrCERT' 에서 'KISA' 탐지
  기타접수/advisory-unknown-002 → 아무 데도 없음 → '-' 자동 지정
PDF 본문에는 카탈로그 제품(Windows·Office·HWP·V3·Chrome…)과 버전 표기(22H2,
점버전, 연도판)를 심어 게이트 자동 추출·색상 하이라이트까지 시연된다.
"""
from __future__ import annotations

import argparse
import json
import shutil
import zipfile
from pathlib import Path

BASE = Path(__file__).resolve().parent.parent
OUT = BASE / "samples" / "demo_advisories"


def _minimal_pdf(lines: list[str]) -> bytes:
    """ASCII 텍스트 1페이지 최소 PDF — app/seed.py 의 헬퍼와 동일 기법(자립 복사본)."""
    def esc(s: str) -> str:
        return "".join(c for c in s if 32 <= ord(c) < 127).replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    content = "BT /F1 11 Tf 50 760 Td 14 TL\n" + "".join(f"({esc(ln)}) Tj T*\n" for ln in lines) + "ET"
    objs = [
        "<</Type/Catalog/Pages 2 0 R>>",
        "<</Type/Pages/Kids[3 0 R]/Count 1>>",
        "<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]/Resources<</Font<</F1 5 0 R>>>>/Contents 4 0 R>>",
        f"<</Length {len(content)}>>\nstream\n{content}\nendstream",
        "<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for i, body in enumerate(objs, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref_pos = len(out)
    out += f"xref\n0 {len(objs)+1}\n0000000000 65535 f \n".encode()
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += f"trailer<</Size {len(objs)+1}/Root 1 0 R>>\nstartxref\n{xref_pos}\n%%EOF".encode()
    return out


# (상대경로, 본문 라인들) — 본문은 ASCII 만 렌더 가능(최소 PDF 폰트 제약).
ADVISORIES: list[tuple[str, list[str]]] = [
    ("국가정보원/2026-07/NIS-2026-0701 MS 정기 보안업데이트 권고.pdf", [
        "National Cyber Security Advisory  2026.7.6",
        "Doc: NIS-2026-0701   Monthly MS Security Update",
        "",
        "Affected: Microsoft Windows 11 22H2, 23H2 / Windows Server 2022",
        "          Microsoft Office 2021 (Excel, Outlook)",
        "",
        "1. CVE-2026-53770  Windows Kernel Elevation of Privilege",
        "2. CVE-2026-53771  Windows SMB Remote Code Execution",
        "3. CVE-2026-53718  Office Excel Memory Corruption",
        "",
        "Apply the July cumulative update immediately.",
    ]),
    ("국가정보원/2026-07/NIS-2026-0702 Windows 커널 취약점 긴급.pdf", [
        "URGENT Security Advisory  2026.7.9",
        "Doc: NIS-2026-0702",
        "",
        "Affected: Microsoft Windows 10 22H2 / Windows 11 23H2",
        "",
        "CVE-2026-53800  Win32k Elevation of Privilege (exploited in the wild)",
        "",
        "Workaround: disable legacy input service until patched.",
    ]),
    ("KISA/KISA-2026-0912 한컴오피스 취약점 조치 권고.pdf", [
        "KrCERT/CC Security Notice  2026.7.3",
        "Doc: KISA-2026-0912   HWP Office Suite",
        "",
        "Affected: HWP 2022, HWP 2024 (Hancom Office)",
        "",
        "1. CVE-2026-61001  HWP postscript filter code execution",
        "2. CVE-2026-61002  HWP OLE object parsing overflow",
        "",
        "Update to the latest Hancom security patch.",
    ]),
    ("KISA/KISA-2026-0913 V3 백신 업데이트 권고.pdf", [
        "KrCERT/CC Security Notice  2026.7.4",
        "Doc: KISA-2026-0913   AhnLab V3",
        "",
        "Affected: AhnLab V3 Lite 4.0, V3 365 Clinic",
        "",
        "CVE-2026-61010  V3 engine local privilege escalation",
        "",
        "Update engine to 2026.07.04.00 or later.",
    ]),
    ("NCSC/웹브라우저/NCSC-2026-0715 Chrome Edge 긴급 패치.pdf", [
        "Browser Emergency Patch Advisory  2026.7.8",
        "Doc: NCSC-2026-0715",
        "",
        "Affected: Google Chrome 126.0.6478.126 and earlier",
        "          Microsoft Edge 126.0.2592.87 and earlier",
        "",
        "1. CVE-2026-62001  Chrome V8 type confusion (0-day)",
        "2. CVE-2026-62002  Edge sandbox escape",
        "",
        "Force-update browsers across all departments.",
    ]),
    ("FSEC-2026-0021 금융권 Java 취약점 권고.pdf", [
        "Financial Sector Security Advisory  2026.7.7",
        "Doc: FSEC-2026-0021",
        "",
        "Affected: Java SE 17.0.11 and earlier (JRE, JDK)",
        "",
        "CVE-2026-63001  JNDI remote class loading vulnerability",
        "",
        "Upgrade to Java SE 17.0.12 / apply vendor mitigation.",
    ]),
    ("기타접수/advisory-unknown-001.pdf", [
        "Forwarded advisory (source in body only)  2026.7.5",
        "Issued by KrCERT vulnerability analysis team.",
        "",
        "Affected: Adobe Acrobat DC 2022",
        "",
        "CVE-2026-64001  Acrobat PDF font parsing use-after-free",
    ]),
    ("기타접수/advisory-unknown-002.pdf", [
        "Anonymous vendor bulletin  2026.7.5",
        "(no issuing organization stated anywhere)",
        "",
        "Affected: OpenSSL 3.0.13 and earlier",
        "",
        "CVE-2026-64002  TLS handshake buffer over-read",
    ]),
]

# 위 PDF 들의 CVE 를 전부 커버하는 내부 형식 피드 — 반입·적용하면 게이트가 해제된다.
FEED = {
    "source": "DEMO-BUNDLE",
    "cves": [
        {"cve_id": "CVE-2026-53770", "product_name": "Microsoft Windows 11", "product_key": "windows_11",
         "affected_versions": ["22H2", "23H2"], "severity": "높음", "cvss_score": 7.8,
         "description": "Windows 커널 권한 상승", "published_at": "2026-07-06", "source": "국가정보원"},
        {"cve_id": "CVE-2026-53771", "product_name": "Microsoft Windows Server", "product_key": "windows_server",
         "affected_versions": ["2022"], "severity": "긴급", "cvss_score": 9.8,
         "description": "SMB 원격 코드 실행", "published_at": "2026-07-06", "source": "국가정보원"},
        {"cve_id": "CVE-2026-53718", "product_name": "Microsoft Office", "product_key": "microsoft_office",
         "affected_versions": ["2021"], "severity": "높음", "cvss_score": 7.8,
         "description": "Excel 메모리 손상", "published_at": "2026-07-06", "source": "국가정보원"},
        {"cve_id": "CVE-2026-53800", "product_name": "Microsoft Windows 10", "product_key": "windows_10",
         "affected_versions": ["22H2"], "severity": "긴급", "cvss_score": 9.0,
         "description": "Win32k 권한 상승(실공격 확인)", "published_at": "2026-07-09", "source": "국가정보원"},
        {"cve_id": "CVE-2026-61001", "product_name": "한컴오피스", "product_key": "hancom_office",
         "affected_versions": ["2022", "2024"], "severity": "높음", "cvss_score": 8.1,
         "description": "HWP 포스트스크립트 필터 코드 실행", "published_at": "2026-07-03", "source": "KISA"},
        {"cve_id": "CVE-2026-61002", "product_name": "한컴오피스", "product_key": "hancom_office",
         "affected_versions": ["2022", "2024"], "severity": "중간", "cvss_score": 6.5,
         "description": "HWP OLE 개체 파싱 오버플로", "published_at": "2026-07-03", "source": "KISA"},
        {"cve_id": "CVE-2026-61010", "product_name": "AhnLab V3", "product_key": "ahnlab_v3",
         "affected_versions": ["4.0"], "severity": "중간", "cvss_score": 6.7,
         "description": "V3 엔진 로컬 권한 상승", "published_at": "2026-07-04", "source": "KISA"},
        {"cve_id": "CVE-2026-62001", "product_name": "Google Chrome", "product_key": "google_chrome",
         "affected_versions": {"lte": "126.0.6478.126"}, "severity": "긴급", "cvss_score": 9.6,
         "description": "V8 타입 혼동(0-day)", "published_at": "2026-07-08", "source": "NCSC"},
        {"cve_id": "CVE-2026-62002", "product_name": "Microsoft Edge", "product_key": "edge",
         "affected_versions": {"lte": "126.0.2592.87"}, "severity": "높음", "cvss_score": 8.8,
         "description": "샌드박스 이탈", "published_at": "2026-07-08", "source": "NCSC"},
        {"cve_id": "CVE-2026-63001", "product_name": "Java SE", "product_key": "java_se",
         "affected_versions": {"lte": "17.0.11"}, "severity": "높음", "cvss_score": 8.1,
         "description": "JNDI 원격 클래스 로딩", "published_at": "2026-07-07", "source": "금융보안원"},
        {"cve_id": "CVE-2026-64001", "product_name": "Adobe Acrobat", "product_key": "adobe_acrobat",
         "affected_versions": ["DC 2022"], "severity": "높음", "cvss_score": 7.8,
         "description": "PDF 폰트 파싱 UAF", "published_at": "2026-07-05", "source": "KISA"},
        {"cve_id": "CVE-2026-64002", "product_name": "OpenSSL", "product_key": "openssl",
         "affected_versions": {"lte": "3.0.13"}, "severity": "중간", "cvss_score": 5.9,
         "description": "TLS 핸드셰이크 버퍼 오버리드", "published_at": "2026-07-05", "source": "-"},
    ],
}


# ── 확장 CVE 피드 — 데모 권고문 CVE 12건을 '포함'하는 슈퍼셋(총 30건) ──
# 영향버전 규칙의 모든 형태(열거 리스트 / lt·lte / range / "*")를 고르게 섞어
# 매칭 엔진·버전 비교기 데모까지 커버한다. 자산대장 샘플의 제품군과 맞물린다.
EXTRA_CVES = [
    {"cve_id": "CVE-2026-70001", "product_name": "Microsoft Windows 10", "product_key": "windows_10",
     "affected_versions": ["21H2", "22H2"], "severity": "높음", "cvss_score": 7.8,
     "description": "그래픽 구성요소 권한 상승", "published_at": "2026-06-10", "source": "국가정보원"},
    {"cve_id": "CVE-2026-70002", "product_name": "Microsoft Windows 11", "product_key": "windows_11",
     "affected_versions": {"lte": "23H2"}, "severity": "긴급", "cvss_score": 9.1,
     "description": "원격 데스크톱 RCE", "published_at": "2026-06-10", "source": "국가정보원"},
    {"cve_id": "CVE-2026-70003", "product_name": "Microsoft Windows Server", "product_key": "windows_server",
     "affected_versions": ["2019", "2022"], "severity": "높음", "cvss_score": 8.1,
     "description": "AD 인증 우회", "published_at": "2026-06-11", "source": "국가정보원"},
    {"cve_id": "CVE-2026-70004", "product_name": "Microsoft Office", "product_key": "microsoft_office",
     "affected_versions": "*", "severity": "중간", "cvss_score": 6.5,
     "description": "매크로 보안 경고 우회(전 버전)", "published_at": "2026-06-12", "source": "NCSC"},
    {"cve_id": "CVE-2026-70005", "product_name": "한컴오피스", "product_key": "hancom_office",
     "affected_versions": ["2020", "2022"], "severity": "중간", "cvss_score": 6.1,
     "description": "HWPX 파서 정보 유출", "published_at": "2026-06-13", "source": "KISA"},
    {"cve_id": "CVE-2026-70006", "product_name": "Google Chrome", "product_key": "google_chrome",
     "affected_versions": {"lt": "125"}, "severity": "높음", "cvss_score": 8.8,
     "description": "WebGPU 힙 오버플로", "published_at": "2026-06-14", "source": "NCTI"},
    {"cve_id": "CVE-2026-70007", "product_name": "Microsoft Edge", "product_key": "edge",
     "affected_versions": {"lt": "126"}, "severity": "중간", "cvss_score": 6.5,
     "description": "SmartScreen 우회", "published_at": "2026-06-14", "source": "NCTI"},
    {"cve_id": "CVE-2026-70008", "product_name": "Java SE", "product_key": "java_se",
     "affected_versions": {"range": ["11", "17.0.11"]}, "severity": "높음", "cvss_score": 7.5,
     "description": "직렬화 역직렬화 RCE", "published_at": "2026-06-15", "source": "금융보안원"},
    {"cve_id": "CVE-2026-70009", "product_name": "Adobe Acrobat", "product_key": "adobe_acrobat",
     "affected_versions": ["DC 2021", "DC 2022"], "severity": "높음", "cvss_score": 7.8,
     "description": "JavaScript API 샌드박스 이탈", "published_at": "2026-06-16", "source": "KISA"},
    {"cve_id": "CVE-2026-70010", "product_name": "OpenSSL", "product_key": "openssl",
     "affected_versions": {"lt": "3.0.14"}, "severity": "높음", "cvss_score": 7.4,
     "description": "세션 재사용 검증 결함", "published_at": "2026-06-17", "source": "NCSC"},
    {"cve_id": "CVE-2026-70011", "product_name": "Apache HTTP Server", "product_key": "apache_httpd",
     "affected_versions": {"lt": "2.4.60"}, "severity": "긴급", "cvss_score": 9.8,
     "description": "mod_proxy 요청 밀반입 RCE", "published_at": "2026-06-18", "source": "NCSC"},
    {"cve_id": "CVE-2026-70012", "product_name": "nginx", "product_key": "nginx",
     "affected_versions": {"lt": "1.25.4"}, "severity": "중간", "cvss_score": 6.5,
     "description": "HTTP/3 QUIC 메모리 누수", "published_at": "2026-06-18", "source": "NCSC"},
    {"cve_id": "CVE-2026-70013", "product_name": "Linux Kernel", "product_key": "linux_kernel",
     "affected_versions": {"lt": "6.8"}, "severity": "높음", "cvss_score": 7.8,
     "description": "netfilter UAF 권한 상승", "published_at": "2026-06-19", "source": "KISA"},
    {"cve_id": "CVE-2026-70014", "product_name": "Ubuntu", "product_key": "ubuntu",
     "affected_versions": {"range": ["20.04", "24.04"]}, "severity": "중간", "cvss_score": 6.4,
     "description": "needrestart 로컬 권한 상승", "published_at": "2026-06-19", "source": "KISA"},
    {"cve_id": "CVE-2026-70015", "product_name": "Microsoft Windows 11", "product_key": "windows_11",
     "affected_versions": ["23H2"], "severity": "중간", "cvss_score": 5.5,
     "description": "탐색기 정보 노출", "published_at": "2026-06-20", "source": "국가정보원"},
    {"cve_id": "CVE-2026-70016", "product_name": "Microsoft Office", "product_key": "microsoft_office",
     "affected_versions": ["2019", "2021"], "severity": "높음", "cvss_score": 7.8,
     "description": "Outlook 미리보기 RCE", "published_at": "2026-06-20", "source": "국가정보원"},
    {"cve_id": "CVE-2026-70017", "product_name": "Google Chrome", "product_key": "google_chrome",
     "affected_versions": {"lte": "126.0.6478.126"}, "severity": "긴급", "cvss_score": 9.6,
     "description": "ANGLE 타입 혼동(실공격)", "published_at": "2026-06-21", "source": "NCTI"},
    {"cve_id": "CVE-2026-70018", "product_name": "Adobe Acrobat", "product_key": "adobe_acrobat",
     "affected_versions": "*", "severity": "낮음", "cvss_score": 3.3,
     "description": "북마크 스크립트 경고 미표시(전 버전)", "published_at": "2026-06-21", "source": "KISA"},
]

# ── 자산관리대장 엣지케이스 샘플 — '복잡한 현장 취합본'을 일부러 재현 ──
# 검증 포인트: 제목행+빈행+2단 병합 헤더 / 부서 세로 병합 / 제품+버전 한 셀 /
# v접두·H차수·'알 수 없음' 버전 / 파일 내 중복 자산번호(최신 행 우선) / 부서·제품·버전
# 누락 경고 / 숫자 자산번호 / 담당자 통합 셀(구분자 분할 시연) / 미매핑 컬럼(extra 보존).
ASSET_TITLE = "2026 상반기 전사 자산관리대장 — 정보보호팀 취합본 (부서 제출 원본 병합)"
ASSET_HEAD_TOP = ["자산번호", "사용부서", "시스템 정보", None, "담당자 정보", None, None, "IP 주소", "구매년도", "비고"]
ASSET_HEAD_SUB = [None, None, "제품/OS", "버전", "이름", "소속팀", "연락처", None, None, None]
ASSET_ROWS: list[list] = [
    ["A-1001", "재무팀", "Windows 11", "23H2", "김영수", "재무1셀", "010-1111-0001", "10.10.1.11", 2023, None],
    ["A-1002", "재무팀", "Windows 11 22H2", None, "이수진", "재무1셀", "010-1111-0002", "10.10.1.12", 2022, "제품 셀에 버전 포함"],
    ["A-1003", "재무팀", "MS Office", "2021", "박지훈", "재무2셀", "010-1111-0003", "10.10.1.13", 2021, None],
    ["A-2001", "개발팀", "Google Chrome", "v124.0.6367.91", "최민아", "플랫폼셀", "010-2222-0001", "10.20.3.21", 2024, "v접두 버전"],
    ["A-2001", "개발팀", "Google Chrome", "v126.0.6478.100", "최민아", "플랫폼셀", "010-2222-0001", "10.20.3.21", 2024, "중복 자산번호 — 이 행(최신)이 남아야 함"],
    ["A-2002", "개발팀", "Chrome 122.x", None, "한도윤", "플랫폼셀", "010-2222-0002", "10.20.3.22", None, None],
    ["A-2003", "개발팀", "Java SE", "17.0.11", "정우석", "백엔드셀", "010-2222-0003", "10.20.3.23", 2020, None],
    ["A-3001", "총무팀", "한글 2022", None, "김보라", "총무셀", "010-3333-0001", "10.30.1.5", 2022, None],
    ["A-3002", "  총무팀  ", "한컴오피스", "2024", "오세훈", "총무셀", "010-3333-0002", "10.30.1.6", 2024, "부서명 앞뒤 공백"],
    ["A-3003", "총무팀", "AhnLab V3 Lite", "4.0", "김보라", "총무셀", "010-3333-0001", "10.30.1.7", None, "공용 PC(담당자 중복)"],
    [20250001, "인사팀", "Microsoft Edge", "126.0.2592.87", "유하늘", "인사셀", "02-555-0100", "10.40.2.31", 2025, "자산번호가 숫자 셀"],
    ["A-5001", "서버운영팀", "OpenSSL", "3.0.13", "백승호", "인프라셀", "010-5555-0001", "10.50.0.10", None, None],
    ["A-5002", "서버운영팀", "Windows Server 2022", None, "백승호", "인프라셀", "010-5555-0001", "10.50.0.0/24", None, "IP 대역 표기"],
    ["A-5003", "서버운영팀", "Apache HTTPD", "2.4.58", "문지영", "인프라셀", "010-5555-0002", "10.50.0.11", None, None],
    [None, None, None, None, None, None, None, None, None, None],   # 중간 빈 행
    ["A-6001", "감사실", "Windows 10 22H2", None, "신재민", "감사셀", "010-6666-0001", "10.60.1.2", 2019, None],
    ["A-6002", "감사실", "Adobe Acrobat", "DC 2022", "신재민 / 감사셀 / 010-6666-0001", None, None, "10.60.1.3", None, "담당자 통합 셀 — '구분자 분할'로 나눠보기"],
    ["A-7001", "미래전략TF", "Ubuntu", "22.04", "장예린", None, None, "10.70.9.9", None, "신설 부서(자동 생성)"],
    ["A-7002", "미래전략TF", "리눅스 커널", "알 수 없음", "장예린", None, None, "10.70.9.10", None, "버전 해석불가 → 후보 매칭"],
    ["A-8001", None, "Windows 11", "23H2", "무명씨", None, None, "10.80.1.1", None, "부서 누락 → 적재 제외 경고"],
    ["A-8002", "품질보증팀", None, None, "나검수", "QA셀", "010-8888-0002", "10.80.1.2", None, "제품 누락 경고"],
    ["A-8003", "품질보증팀", "사내개발 ERP 클라이언트", "7.3.1", "나검수", "QA셀", "010-8888-0002", "10.80.1.3", None, "사전에 없는 제품(슬러그 키)"],
    [None, None, None, None, None, None, None, None, None, "※ 하반기 도입 예정 장비는 별도 시트 참조"],  # 비고만 있는 행
    ["A-9001", "대외협력팀", "Microsoft Office", None, "하은성", None, "02-555-0200", None, None, "버전 누락(적재는 됨 — '*'·전체 규칙 매칭)"],
]

FEED_EXTENDED_NAME = "cve_feed_extended_2026-07.json"
ASSET_XLSX_NAME = "자산관리대장_엣지케이스_샘플.xlsx"


def build_extended_feed() -> Path:
    """확장 피드 — 데모 권고문 피드(12건)를 포함하는 슈퍼셋(총 30건)."""
    dst = BASE / "samples" / FEED_EXTENDED_NAME
    payload = {"source": "DEMO-EXTENDED",
               "cves": FEED["cves"] + EXTRA_CVES}
    dst.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return dst


def build_asset_xlsx() -> Path:
    from openpyxl import Workbook

    dst = BASE / "samples" / ASSET_XLSX_NAME
    wb = Workbook()
    ws = wb.active
    ws.title = "자산대장"
    ws.append([ASSET_TITLE])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws.append([None] * 10)                       # 빈 행(제목과 헤더 분리 — 현장 흔한 형태)
    ws.append(ASSET_HEAD_TOP)                    # 3행: 상단 헤더
    ws.append(ASSET_HEAD_SUB)                    # 4행: 하단 헤더
    ws.merge_cells("C3:D3")                      # 시스템 정보(가로 병합)
    ws.merge_cells("E3:G3")                      # 담당자 정보(가로 병합)
    for col in ("A", "B", "H", "I", "J"):        # 단일 라벨은 세로 병합(2단 헤더)
        ws.merge_cells(f"{col}3:{col}4")
    for row in ASSET_ROWS:
        ws.append(row)
    ws.merge_cells("B5:B7")                      # 재무팀 — 부서 셀 세로 병합(같은 부서 여러 자산)
    wb.save(dst)
    return dst


def build() -> Path:
    if OUT.exists():
        shutil.rmtree(OUT)
    for rel, lines in ADVISORIES:
        path = OUT / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(_minimal_pdf(lines))
    (OUT / "cve_feed_demo.json").write_text(
        json.dumps(FEED, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    readme = OUT / "READ ME — 데모 사용법.txt"
    readme.write_text(
        "보안권고문 데모 샘플 (올인원)\n"
        "================================\n\n"
        "1) 관리자 화면(/admin) → 권고문 처리 → 이 폴더(demo_advisories)를 통째로\n"
        "   업로드 존에 드래그&드롭하거나 '📁 폴더 업로드'로 선택하세요.\n"
        "   - 국가정보원/2026-07/…      → 상위 폴더명에서 출처 자동 탐지\n"
        "   - KISA/…                    → 폴더명에서 탐지\n"
        "   - NCSC/웹브라우저/…         → 상위 폴더명에서 탐지\n"
        "   - FSEC-2026-0021 ….pdf      → 파일명 별칭(FSEC→금융보안원)에서 탐지\n"
        "   - 기타접수/…-001.pdf        → 본문(KrCERT→KISA)에서 탐지\n"
        "   - 기타접수/…-002.pdf        → 미탐지 → '-' 자동 지정(수동 입력 데모)\n"
        "   업로드 직후 '출처 지정' 패널에서 후보 클릭 선택(복수 가능)·일괄/개별 적용을 시험하세요.\n\n"
        "2) 각 권고문 '처리 →' STEP2 에서 DB 미등록 CVE 의 직접 등록 폼과\n"
        "   PDF 색상 하이라이트(제품·버전·날짜 — 폼 필드 색과 동일)를 확인하세요.\n\n"
        "3) 일괄 해제하려면 CVE 데이터베이스 탭에 피드 파일을 드래그&드롭(또는 파일 선택)해\n"
        "   반입·적용하세요.\n"
        "   - cve_feed_demo.json                 : 위 PDF 들의 CVE 12건 (최소 세트)\n"
        "   - samples/cve_feed_extended_2026-07.json : 위 12건을 포함한 확장 30건 —\n"
        "     열거/미만(lt·lte)/범위(range)/전체(*) 등 영향버전 규칙 전 형태 포함\n\n"
        "4) 자산 매칭까지 보려면 samples/자산관리대장_엣지케이스_샘플.xlsx 를\n"
        "   자산관리대장 탭에서 가져오세요 — 헤더 시작행 3, 헤더 행 수 2 로 지정.\n"
        "   제목행·2단 병합 헤더·부서 세로 병합·제품+버전 한 셀·v접두/'알 수 없음' 버전·\n"
        "   파일 내 중복 자산번호·부서/제품/버전 누락·담당자 통합 셀 같은 현장 엣지케이스가\n"
        "   기대대로(경고·최신 행 우선·자동 분리·후보 매칭) 처리되는지 확인하는 용도입니다.\n",
        encoding="utf-8")
    return OUT


def make_zip(zip_path: Path) -> Path:
    zip_path.parent.mkdir(parents=True, exist_ok=True)
    extras = [BASE / "samples" / FEED_EXTENDED_NAME, BASE / "samples" / ASSET_XLSX_NAME]
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob("*")):
            if p.is_file():
                z.write(p, Path("demo_advisories") / p.relative_to(OUT))
        for p in extras:
            if p.exists():
                z.write(p, p.name)
    return zip_path


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--zip", metavar="PATH", help="생성 후 전체를 zip 으로 묶을 경로")
    args = ap.parse_args()
    out = build()
    n = sum(1 for _ in out.rglob("*.pdf"))
    print(f"[demo-samples] {out} — PDF {n}건 + cve_feed_demo.json")
    fx = build_extended_feed()
    print(f"[demo-samples] {fx} — CVE {len(FEED['cves']) + len(EXTRA_CVES)}건(데모 12건 포함 슈퍼셋)")
    ax = build_asset_xlsx()
    print(f"[demo-samples] {ax} — 자산 {sum(1 for r in ASSET_ROWS if any(c not in (None, '') for c in r))}행(엣지케이스)")
    if args.zip:
        zp = make_zip(Path(args.zip))
        print(f"[demo-samples] zip: {zp} ({zp.stat().st_size:,} bytes)")
